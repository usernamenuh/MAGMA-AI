import unittest
import xml.etree.ElementTree as ET
import tempfile
from pathlib import Path

from docx import Document as WordDocument

from utils.annotated_pdf import _page_marks, build_annotated_pdf, count_letters
from utils.docling_extractor import readable_pdf_path


class AnnotatedPdfMatchingTests(unittest.TestCase):
    def test_word_document_can_show_saved_text_in_pdf_preview(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "contoh.docx"
            word = WordDocument()
            word.add_paragraph("Informasi penting sudah tersimpan dalam database.")
            word.add_paragraph("Bagian berikutnya belum dimasukkan ke database.")
            word.save(source)

            with readable_pdf_path(str(source)) as pdf_path:
                preview, stats = build_annotated_pdf(
                    str(pdf_path),
                    [("Informasi penting sudah tersimpan dalam database.", 1)],
                )

            self.assertTrue(preview.startswith(b"%PDF"))
            self.assertGreater(stats["matched_words"], 0)
            self.assertLess(stats["matched_words"], stats["total_words"])

    def test_letter_count_excludes_digits_spaces_and_punctuation(self):
        self.assertEqual(count_letters("A b-2, é!"), 3)

    def test_only_persisted_words_are_marked_in_original_order(self):
        page = ET.fromstring(
            '<page xmlns="http://www.w3.org/1999/xhtml">'
            '<line><word xMin="10" xMax="30" yMax="20">Satu</word>'
            '<word xMin="34" xMax="50" yMax="20">dua</word>'
            '<word xMin="54" xMax="78" yMax="20">tiga</word></line>'
            '</page>'
        )
        marks, matched, total = _page_marks(page, "Satu tiga")
        self.assertEqual((matched, total), (2, 3))
        self.assertEqual(marks, [(10.0, 30.0, 20.0), (54.0, 78.0, 20.0)])

    def test_adjacent_matched_words_share_one_line(self):
        page = ET.fromstring(
            '<page xmlns="http://www.w3.org/1999/xhtml">'
            '<line><word xMin="10" xMax="30" yMax="20">Satu</word>'
            '<word xMin="34" xMax="50" yMax="20">dua</word></line>'
            '</page>'
        )
        marks, matched, total = _page_marks(page, "Satu dua")
        self.assertEqual((matched, total), (2, 2))
        self.assertEqual(marks, [(10.0, 50.0, 20.0)])

    def test_reordered_lines_match_without_exceeding_saved_occurrences(self):
        page = ET.fromstring(
            '<page xmlns="http://www.w3.org/1999/xhtml">'
            '<line><word xMin="10" xMax="20" yMax="20">Alpha</word>'
            '<word xMin="24" xMax="34" yMax="20">Beta</word>'
            '<word xMin="38" xMax="48" yMax="20">Gamma</word></line>'
            '<line><word xMin="10" xMax="20" yMax="40">Delta</word>'
            '<word xMin="24" xMax="34" yMax="40">Epsilon</word>'
            '<word xMin="38" xMax="48" yMax="40">Zeta</word></line>'
            '</page>'
        )
        marks, matched, total = _page_marks(
            page, "Delta Epsilon Zeta Alpha Beta Gamma"
        )
        self.assertEqual((matched, total), (6, 6))
        self.assertEqual(len(marks), 2)

        duplicate = ET.fromstring(
            '<page xmlns="http://www.w3.org/1999/xhtml">'
            '<line><word xMin="10" xMax="20" yMax="20">Alpha</word>'
            '<word xMin="24" xMax="34" yMax="20">Beta</word>'
            '<word xMin="38" xMax="48" yMax="20">Gamma</word></line>'
            '<line><word xMin="10" xMax="20" yMax="40">Alpha</word>'
            '<word xMin="24" xMax="34" yMax="40">Beta</word>'
            '<word xMin="38" xMax="48" yMax="40">Gamma</word></line>'
            '</page>'
        )
        _, matched, total = _page_marks(duplicate, "Alpha Beta Gamma")
        self.assertEqual((matched, total), (3, 6))


if __name__ == "__main__":
    unittest.main()
